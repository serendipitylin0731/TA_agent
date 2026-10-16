"""Build the editable project presentation from screenshots of the running app."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUTPUT = HERE / "知页TA_Agent_项目汇报.pptx"

SW, SH = 13.333, 7.5
NAVY = "142F3C"
NAVY_2 = "203D49"
GREEN = "39795D"
GREEN_2 = "6FA78A"
MINT = "E6F1E9"
PALE = "F5F8F6"
WHITE = "FFFFFF"
INK = "20323D"
MUTED = "668079"
LINE = "D8E5DE"
GOLD = "E5C889"
FONT = "Noto Sans CJK SC"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def inch(value: float):
    return Inches(value)


def rect(slide, x, y, w, h, fill, radius=True, border=None, line_width=0.8):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, inch(x), inch(y), inch(w), inch(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if border:
        shape.line.color.rgb = rgb(border)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    if radius:
        try:
            shape.adjustments[0] = 0.13
        except (AttributeError, IndexError):
            pass
    return shape


def circle(slide, x, y, d, fill, border=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(x), inch(y), inch(d), inch(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if border:
        shape.line.color.rgb = rgb(border)
    else:
        shape.line.fill.background()
    return shape


def txt(
    slide,
    content,
    x,
    y,
    w,
    h,
    size=18,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin=0,
    line_spacing=1.18,
):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = inch(margin)
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = content
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    for run in paragraph.runs:
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def fit_image(slide, path, x, y, w, h, radius=False):
    """Crop without distorting the screenshot and embed the cropped pixels."""
    with Image.open(path) as raw:
        image = raw.convert("RGB")
        iw, ih = image.size
        target = w / h
        current = iw / ih
        if current > target:
            new_w = round(ih * target)
            left = (iw - new_w) // 2
            image = image.crop((left, 0, left + new_w, ih))
        else:
            new_h = round(iw / target)
            top = (ih - new_h) // 2
            image = image.crop((0, top, iw, top + new_h))
        stream = BytesIO()
        image.save(stream, "JPEG", quality=92, optimize=True)
        stream.seek(0)
        pic = slide.shapes.add_picture(stream, inch(x), inch(y), width=inch(w), height=inch(h))
    if radius:
        # A thin frame visually joins the image to its card without modifying the screenshot.
        frame = rect(slide, x, y, w, h, WHITE, radius=True, border=LINE)
        frame.fill.background()
    return pic


def add_logo(slide, x=0.72, y=0.32, light=False):
    rect(slide, x, y, 0.48, 0.48, MINT if light else GREEN, radius=True)
    txt(slide, "知", x + 0.055, y + 0.025, 0.37, 0.4, 19, GREEN if light else WHITE, True, PP_ALIGN.CENTER)
    txt(slide, "知页 TA Agent", x + 0.62, y + 0.05, 2.55, 0.37, 16, WHITE if light else NAVY, True)


def header(slide, eyebrow, title, subtitle=None, dark=False):
    bg = NAVY if dark else PALE
    rect(slide, 0, 0, SW, SH, bg, radius=False)
    add_logo(slide, light=dark)
    txt(slide, eyebrow.upper(), 0.74, 1.08, 11.8, 0.32, 10.5, GREEN_2 if dark else GREEN, True)
    txt(slide, title, 0.72, 1.43, 11.9, 0.57, 28, WHITE if dark else NAVY, True)
    if subtitle:
        txt(slide, subtitle, 0.74, 2.02, 11.8, 0.42, 13.5, "B8D0C7" if dark else MUTED)


def footer(slide, page, dark=False):
    color = "9AB8AB" if dark else MUTED
    rect(slide, 0.72, 7.08, 11.86, 0.01, "416553" if dark else LINE, radius=False)
    txt(slide, "知页 TA Agent · 教学研究原型", 0.74, 7.13, 5, 0.2, 8.5, color)
    txt(slide, f"{page:02d} / 07", 11.6, 7.13, 0.98, 0.2, 8.5, color, align=PP_ALIGN.RIGHT)


def slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, SW, SH, NAVY, radius=False)
    rect(slide, 0.72, 1.38, 0.13, 3.55, GREEN_2, radius=False)
    add_logo(slide, light=True)
    txt(slide, "教材知识库 × 智能问答", 1.05, 1.35, 5.7, 0.5, 16, "A7C9B7", True)
    txt(slide, "知页 TA Agent", 1.03, 2.0, 5.9, 0.95, 38, WHITE, True)
    txt(slide, "让每个教学回答\n都能回到教材", 1.04, 3.02, 5.55, 1.45, 25, WHITE, True, valign=MSO_ANCHOR.TOP, line_spacing=1.25)
    txt(slide, "面向初高中师生的教材检索、问答与阅读平台", 1.05, 4.9, 5.7, 0.54, 14, "B8D0C7")
    rect(slide, 6.94, 1.48, 5.65, 4.40, NAVY_2, border="52746A")
    fit_image(slide, ASSETS / "reader_focus.jpg", 7.12, 1.66, 5.29, 4.05)
    for x, number, caption in [(1.06, "113", "册教材"), (2.73, "20", "个目录位"), (4.42, "3", "个使用入口")]:
        txt(slide, number, x, 6.18, 1.3, 0.45, 25, WHITE, True)
        txt(slide, caption, x, 6.64, 1.45, 0.25, 10.5, "A7C9B7")
    footer(slide, 1, True)


def slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "01 · WHY", "一个教学问题，为什么需要知页？", "从“知道答案”走向“找到教材依据”。")
    items = [
        ("01", "教材分散", "不同学段、学科与册次分散存放，找原文费时。"),
        ("02", "说法不一", "学生的自然提问与教材用语未必一致。"),
        ("03", "答案难核验", "只有结论、没有书名和页码，不便课堂使用。"),
    ]
    for i, (num, title, body) in enumerate(items):
        y = 2.70 + i * 1.07
        rect(slide, 0.74, y, 5.75, 0.88, WHITE, border=LINE)
        circle(slide, 0.96, y + 0.19, 0.47, MINT)
        txt(slide, num, 1.015, y + 0.225, 0.36, 0.18, 10.5, GREEN, True, PP_ALIGN.CENTER)
        txt(slide, title, 1.59, y + 0.12, 4.7, 0.25, 15.5, NAVY, True)
        txt(slide, body, 1.59, y + 0.43, 4.68, 0.32, 10.8, MUTED)
    rect(slide, 6.78, 2.7, 5.78, 3.37, WHITE, border=LINE)
    fit_image(slide, ASSETS / "reader_focus.jpg", 6.94, 2.87, 5.46, 2.65)
    txt(slide, "原版教材可在系统内双页翻阅，直接核对引用。", 7.00, 5.6, 5.35, 0.28, 10.4, MUTED)
    rect(slide, 0.74, 6.28, 11.82, 0.50, MINT)
    txt(slide, "目标路径    自然提问  →  找到教材页  →  整理成答案  →  回到原文核验", 1.00, 6.34, 11.3, 0.34, 14, GREEN, True)
    footer(slide, 2)


def slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "02 · EXPERIENCE", "三个入口，覆盖“问、找、读”", "同一份教材数据，服务不同的学习动作。")
    cards = [
        (0.72, "01  Chat 问答", "跨页整合解释，并标出教材与 PDF 页码。", "chat_focus.jpg"),
        (4.87, "02  查询原文", "模糊查找，按书合并并展示关键页。", "search_focus.jpg"),
        (9.02, "03  阅读教材", "双页翻阅、键盘翻页和命中页跳转。", "reader_focus.jpg"),
    ]
    for x, title, desc, image in cards:
        rect(slide, x, 2.66, 3.60, 3.94, WHITE, border=LINE)
        fit_image(slide, ASSETS / image, x + 0.16, 2.83, 3.28, 2.22)
        txt(slide, title, x + 0.18, 5.22, 3.23, 0.38, 17, NAVY, True)
        txt(slide, desc, x + 0.18, 5.69, 3.23, 0.58, 11.2, MUTED, valign=MSO_ANCHOR.TOP)
    footer(slide, 3)


def slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "03 · HOW", "先检索证据，再生成回答", "知识库仅收录教材；每条引用都能定位到真实 PDF 页。")
    metrics = [("113", "册已接入教材"), ("20", "个学科目录位"), ("10 + 10", "最近对话与检索")]
    for i, (value, label) in enumerate(metrics):
        x = 0.73 + i * 4.11
        rect(slide, x, 2.66, 3.78, 1.08, WHITE, border=LINE)
        txt(slide, value, x + 0.22, 2.80, 1.60, 0.50, 24, GREEN, True)
        txt(slide, label, x + 1.48, 2.85, 2.03, 0.39, 11.2, NAVY, True)
    labels = ["教材 PDF", "分页索引", "多词组检索", "模型整理", "页码引用"]
    for i, label in enumerate(labels):
        x = 0.78 + i * 2.52
        rect(slide, x, 4.26, 2.07, 0.90, MINT if i % 2 == 0 else WHITE, border=LINE)
        circle(slide, x + 0.19, 4.48, 0.39, GREEN)
        txt(slide, str(i + 1), x + 0.275, 4.575, 0.22, 0.18, 10.0, WHITE, True, PP_ALIGN.CENTER)
        txt(slide, label, x + 0.69, 4.47, 1.20, 0.42, 13.4, NAVY, True)
        if i < len(labels) - 1:
            txt(slide, "→", x + 2.13, 4.48, 0.38, 0.42, 17, GREEN, True, PP_ALIGN.CENTER)
    rect(slide, 0.78, 5.64, 11.80, 0.96, WHITE, border=LINE)
    rect(slide, 0.78, 5.64, 0.10, 0.96, GOLD, radius=False)
    txt(slide, "检索不足时", 1.08, 5.77, 2.1, 0.31, 15, NAVY, True)
    txt(slide, "模型仍给出通顺回答；页面在下方单独说明“未检索到合适教材内容”，不编造出处。", 3.05, 5.75, 9.02, 0.55, 11.8, MUTED)
    footer(slide, 4)


def slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "04 · LIVE DEMO", "现在进入网页演示", "点击下方链接，现场体验 Chat 问答、原文检索与教材阅读。", dark=True)
    button = rect(slide, 1.16, 2.79, 11.01, 2.24, NAVY_2, border="648F78", line_width=1.4)
    button.click_action.hyperlink.address = "http://127.0.0.1:8000/"
    circle(slide, 1.62, 3.34, 1.12, GREEN)
    txt(slide, "↗", 1.81, 3.54, 0.75, 0.52, 30, WHITE, True, PP_ALIGN.CENTER)
    txt(slide, "打开知页 TA Agent", 3.07, 3.29, 7.72, 0.55, 24, WHITE, True)
    link = txt(slide, "http://127.0.0.1:8000/", 3.10, 3.96, 7.65, 0.46, 17, "ACD8B9")
    for paragraph in link.text_frame.paragraphs:
        for run in paragraph.runs:
            run.hyperlink.address = "http://127.0.0.1:8000/"
    txt(slide, "本机演示 · 请先在演示电脑启动服务", 1.2, 5.45, 10.9, 0.42, 12.2, "B7CEC2", align=PP_ALIGN.CENTER)
    footer(slide, 5, True)


def slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "05 · VALUE", "优势与便利性", "让教师备课与学生自学都更快、更容易核验。")
    cards = [
        (0.73, 2.69, "01", "有据可查", "答案附教材名称和页码，点击即可回到原文。"),
        (6.78, 2.69, "02", "更容易找到", "自然问句、模糊词和错别字也能进入检索。"),
        (0.73, 4.53, "03", "跨章节整合", "把多页、多册的相关信息整理成连贯解释。"),
        (6.78, 4.53, "04", "学习不中断", "保留最近 10 个会话与 10 次检索，继续查读。"),
    ]
    for x, y, num, title, body in cards:
        rect(slide, x, y, 5.76, 1.55, WHITE, border=LINE)
        circle(slide, x + 0.25, y + 0.26, 0.62, MINT)
        txt(slide, num, x + 0.36, y + 0.38, 0.39, 0.22, 13, GREEN, True, PP_ALIGN.CENTER)
        txt(slide, title, x + 1.09, y + 0.28, 4.23, 0.38, 18.4, NAVY, True)
        txt(slide, body, x + 1.09, y + 0.78, 4.23, 0.50, 12, MUTED, valign=MSO_ANCHOR.TOP)
    footer(slide, 6)


def slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "06 · NEXT", "下一步：让教材知识服务更多教学场景", "在保持页码可核验的前提下，逐步扩大覆盖与能力。")
    cards = [
        (0.72, "01", "扫描页 OCR", "补齐暂无文字层的教材页面，让图中文字也能被检索。"),
        (3.87, "02", "多版本教材", "按授权接入更多版本与校本资料，并保留版本标识。"),
        (7.02, "03", "图文联合检索", "把示意图、表格与相邻正文关联，帮助理解复杂概念。"),
        (10.17, "04", "个性化学习", "形成适配年级和学习目标的复习线索与练习建议。"),
    ]
    for x, num, title, body in cards:
        rect(slide, x, 2.77, 2.46, 3.17, WHITE, border=LINE)
        circle(slide, x + 0.21, 3.08, 0.60, MINT)
        txt(slide, num, x + 0.32, 3.21, 0.38, 0.20, 12, GREEN, True, PP_ALIGN.CENTER)
        txt(slide, title, x + 0.21, 3.97, 2.04, 0.48, 16.6, NAVY, True)
        txt(slide, body, x + 0.21, 4.62, 2.04, 0.92, 11.6, MUTED, valign=MSO_ANCHOR.TOP, line_spacing=1.23)
    txt(slide, "欢迎交流", 0.73, 6.29, 1.4, 0.28, 12, GREEN, True)
    email = txt(slide, "serendipity_lin@sjtu.edu.cn", 2.05, 6.27, 3.95, 0.32, 11.4, NAVY)
    homepage = txt(slide, "serendipitylin0731.github.io", 7.85, 6.27, 4.65, 0.32, 11.4, NAVY)
    for box, address in [(email, "mailto:serendipity_lin@sjtu.edu.cn"), (homepage, "https://serendipitylin0731.github.io/")]:
        for paragraph in box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.hyperlink.address = address
    footer(slide, 7)


def main():
    for image in ["chat_focus.jpg", "search_focus.jpg", "reader_focus.jpg"]:
        if not (ASSETS / image).exists():
            raise FileNotFoundError(ASSETS / image)
    prs = Presentation()
    prs.slide_width = inch(SW)
    prs.slide_height = inch(SH)
    for make_slide in [slide_1, slide_2, slide_3, slide_4, slide_5, slide_6, slide_7]:
        make_slide(prs)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
